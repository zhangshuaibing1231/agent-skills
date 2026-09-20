#!/usr/bin/env python3
import argparse, json
from datetime import date, datetime, timedelta
from pathlib import Path
from docx import Document

CN_NUM = '零壹贰叁肆伍陆柒捌玖'
CN_UNIT = ['', '拾', '佰', '仟']
CN_SEC = ['', '万', '亿']

def section_to_cn(n: int) -> str:
    if n == 0:
        return ''
    s = ''
    unit_pos = 0
    zero = True
    while n > 0:
        d = n % 10
        if d == 0:
            if not zero:
                zero = True
                s = '零' + s
        else:
            zero = False
            s = CN_NUM[d] + CN_UNIT[unit_pos] + s
        unit_pos += 1
        n //= 10
    return s.rstrip('零')

def int_to_cn(n: int) -> str:
    if n == 0:
        return '零'
    parts = []
    sec = 0
    need_zero = False
    while n > 0:
        part = n % 10000
        if part:
            text = section_to_cn(part) + CN_SEC[sec]
            if need_zero and parts and not text.endswith('零'):
                text += '零'
            parts.append(text)
            need_zero = part < 1000
        else:
            need_zero = True
        n //= 10000
        sec += 1
    result = ''.join(reversed(parts)).replace('零零', '零').rstrip('零')
    return result

def money_upper(v):
    f = float(v)
    if not f.is_integer():
        raise ValueError('当前脚本只支持整数元金额')
    return int_to_cn(int(f)) + '元整'

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

def add_one_year(d):
    try:
        return d.replace(year=d.year + 1)
    except ValueError:
        return d.replace(year=d.year + 1, day=28)

def cn_date(d):
    return f'{d.year}年{d.month:02d}月{d.day:02d}日'

def replace_all(doc, mapping):
    for p in doc.paragraphs:
        for k,v in mapping.items():
            if k in p.text:
                for run in p.runs:
                    if k in run.text:
                        run.text = run.text.replace(k, str(v))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text
                if any(k in txt for k in mapping):
                    for p in cell.paragraphs:
                        for run in p.runs:
                            for k,v in mapping.items():
                                if k in run.text:
                                    run.text = run.text.replace(k, str(v))
                    # fallback for placeholders fragmented across runs
                    if any(k in cell.text for k in mapping):
                        new = cell.text
                        for k,v in mapping.items():
                            new = new.replace(k, str(v))
                        cell.text = new

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--data', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--wechat-out')
    args=ap.parse_args()

    data=json.loads(Path(args.data).read_text(encoding='utf-8'))
    confirm=parse_date(data['confirm_date'])
    start=parse_date(data['service_start']) if data.get('service_start') else confirm + timedelta(days=1)
    end=parse_date(data['service_end']) if data.get('service_end') else add_one_year(start) - timedelta(days=1)

    software=int(data.get('software_amount', 998))
    total=int(data['total_amount'])
    hardware_name=(data.get('hardware_name') or '').strip()
    if data.get('hardware_amount') is None:
        hardware=max(0, total-software)
    else:
        hardware=int(data['hardware_amount'])
    if software + hardware != total:
        raise ValueError(f'金额不一致：软件{software}+硬件{hardware}!={total}')
    if not hardware_name and hardware != 0:
        raise ValueError('存在硬件金额但未填写 hardware_name')

    entity_type=data.get('entity_type','company')
    role='法定代表人' if entity_type=='company' else '经营者'
    contact=data['contact'].strip()
    customer=data['customer_name'].strip()
    order_no=data['order_no'].strip()

    mapping={
        '{{ORDER_NO}}':order_no,
        '{{CONFIRM_DATE_CN}}':cn_date(confirm),
        '{{CUSTOMER_NAME}}':customer,
        '{{USCC}}':data['uscc'].strip(),
        '{{CONTACT}}':contact,
        '{{PHONE}}':data['phone'].strip(),
        '{{SOFTWARE_AMOUNT}}':str(software),
        '{{HARDWARE_NAME}}':hardware_name,
        '{{HARDWARE_QTY}}':str(data.get('hardware_qty','1套')) if hardware_name else '',
        '{{HARDWARE_WARRANTY}}':(data.get('hardware_warranty') or '按厂家标准') if hardware_name else '',
        '{{HARDWARE_AMOUNT}}':str(hardware) if hardware_name else '',
        '{{TOTAL_UPPER}}':money_upper(total),
        '{{TOTAL_AMOUNT}}':str(total),
        '{{SERVICE_START_CN}}':cn_date(start),
        '{{SERVICE_END_CN}}':cn_date(end),
        '{{DELIVERY_DATE_CN}}':data.get('delivery_date') or cn_date(start),
        '{{INVOICE_INFO}}':data.get('invoice_info',''),
    }

    doc=Document(args.template)
    replace_all(doc,mapping)
    # If no hardware, blank the hardware row except label.
    if not hardware_name:
        t=doc.tables[1]
        for ci in range(1,5):
            t.cell(2,ci).text=''
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)

    wechat=(f'本人{contact}，为{customer}的{role}。本人已完整阅读并确认《客无忧产品及服务订购确认单》'
            f'（编号：{order_no}）全部内容，同意以本微信作为确认方式，并按照确认单约定订购，订单总金额为人民币{total}元。')
    if args.wechat_out:
        Path(args.wechat_out).write_text(wechat, encoding='utf-8')
    print(wechat)

if __name__=='__main__':
    main()
