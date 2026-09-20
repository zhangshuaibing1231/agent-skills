#!/usr/bin/env python3
import argparse
import sys
from pathlib import Path

from docx import Document


REQUIRED_HEADINGS = (
    "一、本周重点客户情况",
    "二、本周学会/积累了什么销售技能技巧",
    "三、关于客户的常见问题，你积累了哪些",
    "四、行业信息",
    "五、需要公司支持的问题",
)
REQUIRED_SECTIONS = ("跟进中客户", "本周成交客户", "本周卡点客户", "本周流失客户")
FORBIDDEN_TEXT = ("填写说明", "示例客户", "[TODO")


def parse_args():
    parser = argparse.ArgumentParser(description="Validate a generated Kewayou weekly report DOCX.")
    parser.add_argument("--input", required=True, help="Generated DOCX path")
    return parser.parse_args()


def main():
    args = parse_args()
    path = Path(args.input).resolve()
    if not path.exists():
        print(f"ERROR: file not found: {path}", file=sys.stderr)
        return 2

    try:
        document = Document(path)
    except Exception as exc:
        print(f"ERROR: cannot open DOCX: {exc}", file=sys.stderr)
        return 2

    paragraph_text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    all_text = paragraph_text + "\n" + table_text
    errors = []

    if len(document.tables) != 4:
        errors.append(f"expected 4 tables, found {len(document.tables)}")
    for heading in REQUIRED_HEADINGS:
        if heading not in paragraph_text:
            errors.append(f"missing heading: {heading}")
    for section in REQUIRED_SECTIONS:
        if section not in table_text:
            errors.append(f"missing customer section: {section}")
    for value in FORBIDDEN_TEXT:
        if value in all_text:
            errors.append(f"template placeholder remains: {value}")
    if "周报周期：" not in paragraph_text:
        errors.append("missing report period")
    if not document.paragraphs or not document.paragraphs[0].text.strip():
        errors.append("missing report title or author")

    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(f"OK: {path}")
    print(f"tables={len(document.tables)} rows={[len(table.rows) for table in document.tables]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
