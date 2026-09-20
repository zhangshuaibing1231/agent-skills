#!/usr/bin/env python3
import argparse
import json
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SKILL_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_ROOT / "assets" / "weekly_report_template.docx"
REQUIRED_LISTS = (
    "follow_up", "closed_won", "blocked", "lost",
    "skills", "faqs", "industry_info", "support",
)


def parse_args():
    parser = argparse.ArgumentParser(description="Generate a Kewayou weekly report DOCX.")
    parser.add_argument("--input", required=True, help="Structured weekly report JSON")
    parser.add_argument("--output", required=True, help="Output DOCX path")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Optional template DOCX")
    parser.add_argument(
        "--continuation-after", type=int, default=3,
        help="Insert a second blocked-customer header after this many rows; use 0 to disable",
    )
    return parser.parse_args()


def require_fields(item, fields, label):
    if not isinstance(item, dict):
        raise ValueError(f"{label} must contain objects")
    missing = [field for field in fields if field not in item]
    if missing:
        raise ValueError(f"{label} item is missing fields: {', '.join(missing)}")


def load_data(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Top-level JSON must be an object")
    for key in ("author", "period", *REQUIRED_LISTS):
        if key not in data:
            raise ValueError(f"Missing top-level field: {key}")
    if not isinstance(data["period"], dict):
        raise ValueError("period must be an object")
    for field in ("start", "end"):
        if field not in data["period"]:
            raise ValueError(f"period is missing: {field}")
        date.fromisoformat(data["period"][field])
    for key in REQUIRED_LISTS:
        if not isinstance(data[key], list):
            raise ValueError(f"{key} must be an array")

    schemas = {
        "follow_up": ("name", "type", "stage", "progress", "next_action"),
        "closed_won": ("name", "type", "product", "key_reason", "next_action"),
        "blocked": ("name", "type", "blocker", "next_action", "risk"),
        "lost": ("name", "product", "reason", "review"),
        "skills": ("scene", "technique"),
        "faqs": ("question", "concern", "answer"),
        "support": ("issue", "support_needed", "owner"),
    }
    for key, fields in schemas.items():
        for item in data[key]:
            require_fields(item, fields, key)
    for item in data["industry_info"]:
        if not isinstance(item, str):
            raise ValueError("industry_info must contain strings")
    for item in data["blocked"]:
        if item["risk"] not in ("高", "中", "低"):
            raise ValueError("blocked risk must be 高, 中, or 低")
    return data


def set_run_font(run, size=None, bold=None):
    run.font.name = "Microsoft YaHei"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), "Microsoft YaHei")


def set_paragraph_text(paragraph, text, size=None, bold=None, center=None):
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = str(text)
        for run in list(paragraph.runs[1:]):
            paragraph._p.remove(run._r)
    else:
        first = paragraph.add_run(str(text))
    set_run_font(first, size=size, bold=bold)
    if center is not None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08


def set_cell_text(cell, text, size=9.0, bold=None, center=False):
    paragraph = cell.paragraphs[0]
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)
    set_paragraph_text(paragraph, text, size=size, bold=bold, center=center)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def rebuild_table(table, header_tr, body_tr, rows):
    for row in list(table.rows):
        table._tbl.remove(row._tr)
    table._tbl.append(deepcopy(header_tr))
    repeat_header(table.rows[0])
    for values in rows:
        table._tbl.append(deepcopy(body_tr))
        row = table.rows[-1]
        keep_row_together(row)
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value)


def date_label(value):
    parsed = date.fromisoformat(value)
    return f"{parsed.year}年{parsed.month}月{parsed.day}日"


def section_rows(data):
    lost = data["lost"] or [{
        "name": "本周暂无明确流失客户",
        "product": "-",
        "reason": "本周记录中没有客户明确终止合作或转向其他系统。",
        "review": "继续区分暂缓、考虑中和明确流失，避免过早关闭商机。",
    }]
    return {
        "follow_up": [
            (x["name"], x["type"], x["stage"], x["progress"], x["next_action"])
            for x in data["follow_up"]
        ],
        "closed_won": [
            (x["name"], x["type"], x["product"], x["key_reason"], x["next_action"])
            for x in data["closed_won"]
        ],
        "blocked": [
            (x["name"], x["type"], x["blocker"], x["next_action"], x["risk"])
            for x in data["blocked"]
        ],
        "lost": [(x["name"], x["product"], x["reason"], x["review"], "") for x in lost],
    }


def build_report(data, template_path, output_path, continuation_after):
    document = Document(template_path)
    author = str(data["author"]).strip() or "周报"
    start = date_label(data["period"]["start"])
    end = date_label(data["period"]["end"])

    set_paragraph_text(document.paragraphs[0], author, size=22, bold=True)
    period_run = document.paragraphs[0].add_run(f"\n周报周期：{start}—{end}")
    set_run_font(period_run, size=10.5, bold=False)
    set_paragraph_text(document.paragraphs[3], "")

    for index in (19, 18, 17, 16, 15, 9, 5, 2):
        paragraph = document.paragraphs[index]
        paragraph._element.getparent().remove(paragraph._element)

    customer_table = document.tables[0]
    original = [deepcopy(row._tr) for row in customer_table.rows]
    rows = section_rows(data)
    sections = [
        (original[0], original[1], original[2], rows["follow_up"]),
        (original[4], original[5], original[6], rows["closed_won"]),
        (original[8], original[9], original[10], rows["blocked"]),
        (original[12], original[13], original[14], rows["lost"]),
    ]
    for row in list(customer_table.rows):
        customer_table._tbl.remove(row._tr)
    for section_tr, header_tr, body_tr, values_list in sections:
        customer_table._tbl.append(deepcopy(section_tr))
        keep_row_together(customer_table.rows[-1])
        customer_table._tbl.append(deepcopy(header_tr))
        repeat_header(customer_table.rows[-1])
        for values in values_list:
            customer_table._tbl.append(deepcopy(body_tr))
            row = customer_table.rows[-1]
            keep_row_together(row)
            for index, value in enumerate(values):
                set_cell_text(row.cells[index], value)

    if continuation_after > 0 and len(rows["blocked"]) > continuation_after:
        blocked_index = next(
            index for index, row in enumerate(customer_table.rows)
            if row.cells[0].text.strip() == "本周卡点客户"
        )
        header_tr = deepcopy(customer_table.rows[blocked_index + 1]._tr)
        insert_index = blocked_index + 2 + continuation_after
        customer_table.rows[insert_index]._tr.addprevious(header_tr)
        continuation_header = customer_table.rows[insert_index]
        for cell in continuation_header.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
        if insert_index + 1 < len(customer_table.rows):
            for cell in customer_table.rows[insert_index + 1].cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True

    skills_table = document.tables[1]
    rebuild_table(
        skills_table, deepcopy(skills_table.rows[0]._tr), deepcopy(skills_table.rows[1]._tr),
        [(x["scene"], x["technique"]) for x in data["skills"]],
    )

    faq_table = document.tables[2]
    rebuild_table(
        faq_table, deepcopy(faq_table.rows[0]._tr), deepcopy(faq_table.rows[1]._tr),
        [(x["question"], x["concern"], x["answer"]) for x in data["faqs"]],
    )

    industry_paragraphs = [document.paragraphs[9], document.paragraphs[10], document.paragraphs[11]]
    industry_values = data["industry_info"][:3]
    while len(industry_values) < 3:
        industry_values.append("")
    for index, (paragraph, value) in enumerate(zip(industry_paragraphs, industry_values), start=1):
        set_paragraph_text(paragraph, f"{index}. {value}" if value else "", size=10.5)

    support_table = document.tables[3]
    rebuild_table(
        support_table, deepcopy(support_table.rows[0]._tr), deepcopy(support_table.rows[1]._tr),
        [(x["issue"], x["support_needed"], x["owner"]) for x in data["support"]],
    )

    support_heading = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip() == "五、需要公司支持的问题"
    )
    support_heading.paragraph_format.page_break_before = True
    support_heading.paragraph_format.keep_with_next = True

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        set_run_font(run)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main():
    args = parse_args()
    data = load_data(args.input)
    template_path = Path(args.template).resolve()
    output_path = Path(args.output).resolve()
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")
    build_report(data, template_path, output_path, args.continuation_after)
    print(output_path)


if __name__ == "__main__":
    main()
